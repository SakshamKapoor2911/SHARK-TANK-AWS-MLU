import os
import sys
import argparse
import subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import qrcode
from PIL import Image

# ==============================================================================
# THEME: MATRAIX-CALIBRATED EXECUTIVE SLATE (Research & Venture Grade)
# ==============================================================================
BG_COLOR = RGBColor(11, 14, 20)          # #0B0E14 Deep Navy-Charcoal Canvas
SURFACE_COLOR = RGBColor(19, 24, 34)     # #131822 Primary Elevated Card
SURFACE_ALT = RGBColor(26, 32, 46)       # #1A202E Highlighted Card
BORDER_COLOR = RGBColor(43, 60, 85)      # #2B3C55 Subtle Architectural Border

BLUE_ACCENT = RGBColor(74, 143, 194)     # #4A8FC2 Phosphor Blue (Executive Primary)
CYAN_ACCENT = RGBColor(0, 196, 159)      # #00C49F Mint Emerald (Metrics & Success)
PURPLE_ACCENT = RGBColor(168, 85, 247)   # #A855F7 Spectral Violet
AMBER_ACCENT = RGBColor(245, 158, 11)    # #F59E0B Warm Gold (Traps & Urgency)
CORAL_ACCENT = RGBColor(239, 68, 68)     # #EF4444 Contrast Alert

TEXT_WHITE = RGBColor(255, 255, 255)     # Headers
TEXT_LIGHT = RGBColor(232, 233, 236)     # Slate 100
TEXT_MUTED = RGBColor(160, 168, 180)     # High-contrast readable secondary
TEXT_DIM = RGBColor(120, 128, 142)       # Subtle captions

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_PPTX = BASE_DIR / "Prismatic_Shark_Tank_Pitch_Deck.pptx"
OUTPUT_PDF = BASE_DIR / "Prismatic_Shark_Tank_Pitch_Deck.pdf"
ASSETS_DIR = BASE_DIR / "assets"
SLIDES_DIR = ASSETS_DIR / "slides"
GITHUB_QR = ASSETS_DIR / "prismatic_github_qr.png"

def ensure_assets():
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=10,
        border=2
    )
    qr.add_data("https://github.com/SakshamKapoor2911/SHARK-TANK-AWS-MLU")
    qr.make(fit=True)
    # Themed QR code with Mint Cyan modules on Deep Slate background (identical to MatrAIx)
    img = qr.make_image(fill_color=(0, 196, 159), back_color=(11, 14, 20)).convert("RGBA")
    img = img.resize((400, 400), Image.Resampling.LANCZOS)
    img.save(str(GITHUB_QR), "PNG")

def set_slide_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return bg

def add_card(slide, left, top, width, height, bg_color=SURFACE_COLOR, border_color=BORDER_COLOR, border_width=1.0):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_clean_textbox(slide, left, top, width, height, text="", font_size=11, font_color=TEXT_WHITE, bold=False, italic=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    if text:
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Segoe UI"
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.italic = italic
        p.font.color.rgb = font_color
        p.alignment = align
    return tb

def add_header(slide, tracker_tag, title_text, subtitle_text="", timing_badge=""):
    add_clean_textbox(slide, Inches(0.8), Inches(0.38), Inches(9.0), Inches(0.32),
                      text=tracker_tag.upper(), font_size=11.5, font_color=BLUE_ACCENT, bold=True)

    if timing_badge:
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.333), Inches(0.35), Inches(2.2), Inches(0.38))
        pill.fill.solid()
        pill.fill.fore_color.rgb = SURFACE_ALT
        pill.line.color.rgb = BLUE_ACCENT
        pill.line.width = Pt(1.0)
        tf_p = pill.text_frame
        tf_p.margin_left = tf_p.margin_right = tf_p.margin_top = tf_p.margin_bottom = 0
        p_pill = tf_p.paragraphs[0]
        p_pill.text = timing_badge.upper()
        p_pill.font.name = "Segoe UI"
        p_pill.font.size = Pt(10)
        p_pill.font.bold = True
        p_pill.font.color.rgb = CYAN_ACCENT
        p_pill.alignment = PP_ALIGN.CENTER

    add_clean_textbox(slide, Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.55),
                      text=title_text, font_size=23, font_color=TEXT_WHITE, bold=True)

    if subtitle_text:
        add_clean_textbox(slide, Inches(0.8), Inches(1.30), Inches(11.733), Inches(0.35),
                          text=subtitle_text, font_size=12, font_color=TEXT_MUTED)

def add_slide_footer(slide, current_idx, total_slides=4):
    footer_text = f"PRISMATIC · AWS × MLU FALL SYMPOSIUM 2026 · HOWARD UNIVERSITY & AMAZON HQ2 | Slide {current_idx} of {total_slides}"
    add_clean_textbox(slide, Inches(0.8), Inches(7.12), Inches(11.733), Inches(0.25),
                      text=footer_text, font_size=9, font_color=TEXT_DIM)

def build_prismatic_masterpiece():
    ensure_assets()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ==============================================================================
    # SLIDE 1: MASTER TITLE & EXECUTIVE HOOK (SEGMENTS 1 & 2)
    # ==============================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Event Badge Pill
    badge = add_card(s1, Inches(0.8), Inches(0.62), Inches(7.5), Inches(0.42), bg_color=SURFACE_ALT, border_color=BLUE_ACCENT, border_width=1.0)
    add_clean_textbox(s1, Inches(0.95), Inches(0.68), Inches(7.2), Inches(0.30),
                      text="AWS-MLU AI TEACHING & RESEARCH SYMPOSIUM 2026 · HOWARD UNIVERSITY & AMAZON HQ2", 
                      font_size=10, font_color=BLUE_ACCENT, bold=True)

    # Timing Pill Top Right
    add_card(s1, Inches(10.533), Inches(0.62), Inches(2.0), Inches(0.42), bg_color=SURFACE_ALT, border_color=CYAN_ACCENT, border_width=1.0)
    add_clean_textbox(s1, Inches(10.533), Inches(0.68), Inches(2.0), Inches(0.30),
                      text="35s THE HOOK", font_size=10, font_color=CYAN_ACCENT, bold=True, align=PP_ALIGN.CENTER)

    # Big Title Block
    tb_hero = s1.shapes.add_textbox(Inches(0.8), Inches(1.22), Inches(9.2), Inches(1.25))
    tf_h = tb_hero.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
    p_h1 = tf_h.paragraphs[0]
    p_h1.text = "Prismatic: The Adaptive Cognitive Engine"
    p_h1.font.name = "Segoe UI"
    p_h1.font.size = Pt(32)
    p_h1.font.bold = True
    p_h1.font.color.rgb = TEXT_WHITE

    p_h2 = tf_h.add_paragraph()
    p_h2.text = "Refracting Complex Knowledge into 4 Personalized Learning Modalities on Amazon Quick"
    p_h2.font.name = "Segoe UI"
    p_h2.font.size = Pt(14)
    p_h2.font.color.rgb = BLUE_ACCENT
    p_h2.space_before = Pt(4)

    # Presenter & Cohort Card
    p_card = add_card(s1, Inches(0.8), Inches(2.65), Inches(9.0), Inches(2.2), SURFACE_COLOR, BORDER_COLOR)
    tb_pres = s1.shapes.add_textbox(Inches(1.05), Inches(2.80), Inches(8.5), Inches(1.9))
    tf_pres = tb_pres.text_frame
    tf_pres.word_wrap = True
    tf_pres.margin_left = tf_pres.margin_top = tf_pres.margin_right = tf_pres.margin_bottom = 0

    pp1 = tf_pres.paragraphs[0]
    pp1.text = "FOUNDING TEAM & AFFILIATION:"
    pp1.font.name = "Segoe UI"
    pp1.font.size = Pt(12)
    pp1.font.bold = True
    pp1.font.color.rgb = CYAN_ACCENT

    pp2 = tf_pres.add_paragraph()
    r_name = pp2.add_run()
    r_name.text = "Saksham Kapoor "
    r_name.font.bold = True
    r_name.font.size = Pt(13)
    r_name.font.color.rgb = TEXT_WHITE
    r_affil = pp2.add_run()
    r_affil.text = "(University of Maryland, College Park) · 10-Student Cross-Campus Builder Cohort"
    r_affil.font.size = Pt(13)
    r_affil.font.color.rgb = TEXT_MUTED
    pp2.space_after = Pt(4)

    pp3 = tf_pres.add_paragraph()
    pp3.text = "THE 300-STUDENT LECTURE HALL PARADOX:"
    pp3.font.name = "Segoe UI"
    pp3.font.size = Pt(12)
    pp3.font.bold = True
    pp3.font.color.rgb = BLUE_ACCENT

    pp4 = tf_pres.add_paragraph()
    pp4.text = "One professor teaches 300 students in one rigid format. White light contains every color, but a flat chalkboard only projects one shade. When students freeze with intense brain fog, it's not a lack of intelligence—it's a cognitive mismatch."
    pp4.font.name = "Segoe UI"
    pp4.font.size = Pt(11.5)
    pp4.font.color.rgb = TEXT_LIGHT

    # QR Code Card Top Right
    add_card(s1, Inches(10.15), Inches(1.30), Inches(2.383), Inches(3.55), bg_color=SURFACE_COLOR, border_color=BORDER_COLOR)
    if GITHUB_QR.exists():
        s1.shapes.add_picture(str(GITHUB_QR), Inches(10.45), Inches(1.50), Inches(1.78), Inches(1.78))
    
    tb_qr = add_clean_textbox(s1, Inches(10.35), Inches(3.40), Inches(1.98), Inches(1.3))
    tf_qr = tb_qr.text_frame
    pq1 = tf_qr.paragraphs[0]
    pq1.text = "SCAN FOR DEMO"
    pq1.font.name = "Segoe UI"
    pq1.font.size = Pt(11)
    pq1.font.bold = True
    pq1.font.color.rgb = CYAN_ACCENT
    pq1.alignment = PP_ALIGN.CENTER
    pq2 = tf_qr.add_paragraph()
    pq2.text = "GitHub repo, web app code, prompt specs & cue cards."
    pq2.font.name = "Segoe UI"
    pq2.font.size = Pt(9)
    pq2.font.color.rgb = TEXT_MUTED
    pq2.alignment = PP_ALIGN.CENTER

    # 4 Bottom Metric Cards (Matching MatrAIx proportions)
    metrics_s1 = [
        ("< 30 SECONDS", "Cognitive Triage", "5-question diagnostic maps intake profile"),
        ("4 SPECTRUMS", "Learning Modalities", "Analogies, Schemas, Socratic & Decks"),
        ("2-SENTENCE CAP", "Socratic Sparring", "Active recall without text dumps"),
        ("100% NO-CODE", "Amazon Quick", "Full multi-widget pipeline in 60 mins")
    ]
    card_w = Inches(2.76)
    card_gap = Inches(0.23)
    start_x = Inches(0.8)
    top_y = Inches(5.15)
    card_h = Inches(1.75)

    for idx, (m_val, m_lbl, m_desc) in enumerate(metrics_s1):
        cx = start_x + idx * (card_w + card_gap)
        add_card(s1, cx, top_y, card_w, card_h, SURFACE_ALT, BORDER_COLOR)
        
        add_clean_textbox(s1, cx + Inches(0.18), top_y + Inches(0.18), card_w - Inches(0.36), Inches(0.42),
                          text=m_val, font_size=18 if idx == 3 else 20, font_color=CYAN_ACCENT if idx != 3 else PURPLE_ACCENT, bold=True)
        
        add_clean_textbox(s1, cx + Inches(0.18), top_y + Inches(0.64), card_w - Inches(0.36), Inches(0.32),
                          text=m_lbl, font_size=13, font_color=TEXT_WHITE, bold=True)

        add_clean_textbox(s1, cx + Inches(0.18), top_y + Inches(1.02), card_w - Inches(0.36), Inches(0.60),
                          text=m_desc, font_size=11, font_color=TEXT_MUTED)

    add_slide_footer(s1, 1, 4)

    notes1 = s1.notes_slide.notes_text_frame
    notes1.text = (
        "OFFICIAL TIMED SCRIPT (0:00 - 0:35):\n\n"
        "[0:00 - 0:10] SEGMENT 1: TEAM NAME (10s)\n"
        "\"Good afternoon, judges and fellow builders! We are PRISMATIC—and we are here to refract the way college students learn!\"\n\n"
        "[0:10 - 0:35] SEGMENT 2: THE HOOK (25s)\n"
        "\"Picture a 300-person lecture hall right here at Howard or UMD. The professor spends 75 minutes covering one complex concept. Half the class nods; the other half walks out completely lost with intense brain fog.\n"
        "The problem isn't student intellect. White light contains every color, but a flat chalkboard only projects one shade. Professors teach in one rigid format, while human brains process knowledge in completely different spectrums.\""
    )

    # ==============================================================================
    # SLIDE 2: THE 4 REFRACTED LEARNING SPECTRUMS (45s OVERVIEW)
    # ==============================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Product Architecture & Execution", 
               "The 4 Learning Spectrums: Zero Errors, Pure Retention", 
               "Orchestrated in a 60-minute build sprint on Amazon Quick & Amazon Bedrock foundation models", 
               "45s APP OVERVIEW")

    spectrums_data = [
        ("SPECTRUM 1: INTUITIVE ANALOGY", "Dorm & Dining Hall Realities", 
         "Translates abstract theory into daily campus navigation. E.g., Dijkstra's shortest path = navigating peak dining hall line rush at Blackburn Center.", 
         "METAPHORICAL", BLUE_ACCENT),
        ("SPECTRUM 2: VISUAL SCHEMATIC", "Structural Concept Maps & ASCII", 
         "Renders execution trees, logical branches, and state machines into structural mental schemas for complex STEM and CS coursework.", 
         "VISUO-SPATIAL", PURPLE_ACCENT),
        ("SPECTRUM 3: SOCRATIC SPARRING", "Active Retrieval Partner", 
         "Chatbot strictly capped at 2 sentences; never lectures; fires progressive edge cases to force active recall without giving away answers.", 
         "KINESTHETIC", CYAN_ACCENT),
        ("SPECTRUM 4: EXAM TRAP DECK", "4-Slide Micro-Presentation", 
         "Distills the topic into 4 high-yield slides exposing the exact trick questions professors test on Midterm #2, paired with AI visual anchors.", 
         "SYNTHESIS", AMBER_ACCENT)
    ]

    sp_w = Inches(5.72)
    sp_h = Inches(2.45)
    col_gap = Inches(0.28)
    row_gap = Inches(0.25)
    start_top = Inches(1.85)

    for i, (spec_tag, spec_title, spec_desc, badge_text, color) in enumerate(spectrums_data):
        col = i % 2
        row = i // 2
        left_pos = Inches(0.8) + col * (sp_w + col_gap)
        top_pos2 = start_top + row * (sp_h + row_gap)

        add_card(s2, left_pos, top_pos2, sp_w, sp_h, bg_color=SURFACE_COLOR, border_color=BORDER_COLOR)

        # Internal Pill Badge
        add_card(s2, left_pos + Inches(0.25), top_pos2 + Inches(0.22), Inches(1.8), Inches(0.30), bg_color=SURFACE_ALT, border_color=color)
        add_clean_textbox(s2, left_pos + Inches(0.25), top_pos2 + Inches(0.25), Inches(1.8), Inches(0.25),
                          text=badge_text, font_size=8.5, font_color=color, bold=True, align=PP_ALIGN.CENTER)

        add_clean_textbox(s2, left_pos + Inches(0.25), top_pos2 + Inches(0.62), sp_w - Inches(0.5), Inches(0.35),
                          text=spec_title, font_size=14, font_color=TEXT_WHITE, bold=True)

        add_clean_textbox(s2, left_pos + Inches(0.25), top_pos2 + Inches(0.98), sp_w - Inches(0.5), Inches(0.28),
                          text=spec_tag, font_size=10.5, font_color=color, bold=True)

        add_clean_textbox(s2, left_pos + Inches(0.25), top_pos2 + Inches(1.32), sp_w - Inches(0.5), Inches(0.98),
                          text=spec_desc, font_size=11, font_color=TEXT_MUTED)

    add_slide_footer(s2, 2, 4)

    notes2 = s2.notes_slide.notes_text_frame
    notes2.text = (
        "OFFICIAL TIMED SCRIPT (0:35 - 1:20):\n\n"
        "[0:35 - 1:20] SEGMENT 3: YOUR APP (45s)\n"
        "\"We built Prismatic on Amazon Quick—an agentic cognitive engine that turns 1 static lecture into 4 tailored learning modalities in under 30 seconds.\n"
        "First, a rapid 5-question diagnostic maps how your brain absorbs information. Then, our Amazon Bedrock pipeline instantly refracts the material:\n"
        "- Spectrum 1 translates abstract theory into relatable dorm and dining hall analogies.\n"
        "- Spectrum 2 renders structural ASCII concept maps.\n"
        "- Spectrum 3 launches a Socratic Sparring Partner that quizzes you one edge-case at a time without spoiling answers.\n"
        "- And Spectrum 4 builds a 4-slide micro-deck exposing the exact trap questions professors test on midterms.\""
    )

    # ==============================================================================
    # SLIDE 3: WHY IT WINS & COMPETITIVE MOAT (25s)
    # ==============================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Competitive Moat & Student Value", 
               "Why Prismatic Wins: The $80/Hour Tutoring Alternative", 
               "Comparing modern AI study workflows against passive text dumps and expensive human tutoring", 
               "25s WHY IT WINS")

    wins_data = [
        ("VS. CHATGPT & RAW LLMS", "Active Sparring vs. Text Dumps", 
         "ChatGPT dumps 800-word walls students glaze over. Prismatic caps replies at 2 sentences and forces active recall.", 
         BLUE_ACCENT),
        ("VS. $80/HR PRIVATE TUTORING", "Democratizing Academic Mastery", 
         "Tutors cost $80/hr; office hours are 2 hrs/week. Prismatic gives every student 24/7 adaptive tutoring for pennies.", 
         CORAL_ACCENT),
        ("VS. STATIC SLIDES & SYLLABI", "Dynamic Cognitive Fit", 
         "Ends task paralysis for ADHD, visual, and verbal learners with a 30-second diagnostic.", 
         CYAN_ACCENT)
    ]

    w_w = Inches(3.73)
    w_h = Inches(4.9)
    top_pos_w = Inches(1.85)

    for i, (w_tag, w_title, w_desc, color) in enumerate(wins_data):
        left_pos = Inches(0.8) + i * (w_w + Inches(0.27))
        add_card(s3, left_pos, top_pos_w, w_w, w_h, bg_color=SURFACE_COLOR, border_color=BORDER_COLOR)

        add_clean_textbox(s3, left_pos + Inches(0.25), top_pos_w + Inches(0.35), w_w - Inches(0.5), Inches(0.35),
                          text=w_tag, font_size=11.5, font_color=color, bold=True)

        add_clean_textbox(s3, left_pos + Inches(0.25), top_pos_w + Inches(0.75), w_w - Inches(0.5), Inches(0.50),
                          text=w_title, font_size=14, font_color=TEXT_WHITE, bold=True)

        add_clean_textbox(s3, left_pos + Inches(0.25), top_pos_w + Inches(1.35), w_w - Inches(0.5), Inches(3.1),
                          text=w_desc, font_size=12, font_color=TEXT_MUTED)

    add_slide_footer(s3, 3, 4)

    notes3 = s3.notes_slide.notes_text_frame
    notes3.text = (
        "OFFICIAL TIMED SCRIPT (1:20 - 1:45):\n\n"
        "[1:20 - 1:45] SEGMENT 4: WHY IT WINS (25s)\n"
        "\"Over 20 million college students face academic paralysis every semester.\n"
        "Generic tools like ChatGPT just dump dry walls of text that students glaze over.\n"
        "Private tutoring costs $80 an hour, and professor office hours are only 2 hours a week.\n"
        "Prismatic is active, multi-modal, and adapts the syllabus to the student—democratizing elite private tutoring for pennies.\""
    )

    # ==============================================================================
    # SLIDE 4: THE CLOSE & SHARK CRITERIA AUDIT (15s)
    # ==============================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Closing Statement & Evaluation Criteria", 
               "'With More Time We'd Add...' & The Shark Tank Verdict", 
               "Meeting 100% of official judging criteria: Creativity, Usefulness, Execution & Teamwork", 
               "15s THE CLOSE")

    # 4 Judging Criteria Cards
    criteria_data = [
        ("CREATIVITY", "Cognitive refraction metaphor (white light → spectrums)", BLUE_ACCENT),
        ("USEFULNESS", "Directly solves dorm paralysis & exam anxiety", CYAN_ACCENT),
        ("EXECUTION", "Working 6-widget pipeline on Amazon Quick", PURPLE_ACCENT),
        ("TEAMWORK", "10 teammates with dedicated active roles", AMBER_ACCENT)
    ]
    sc_w = Inches(2.76)
    sc_h = Inches(1.55)
    for i, (ct, cd, cc) in enumerate(criteria_data):
        left_pos = Inches(0.8) + i * (sc_w + Inches(0.23))
        add_card(s4, left_pos, Inches(1.85), sc_w, sc_h, bg_color=SURFACE_COLOR, border_color=BORDER_COLOR)

        add_clean_textbox(s4, left_pos + Inches(0.20), Inches(2.00), sc_w - Inches(0.4), Inches(0.32),
                          text=ct, font_size=11.5, font_color=cc, bold=True)

        add_clean_textbox(s4, left_pos + Inches(0.20), Inches(2.38), sc_w - Inches(0.4), Inches(0.90),
                          text=cd, font_size=10.5, font_color=TEXT_LIGHT)

    # Master Close Banner
    banner = add_card(s4, Inches(0.8), Inches(3.65), Inches(11.733), Inches(3.25), bg_color=SURFACE_ALT, border_color=BLUE_ACCENT, border_width=1.2)
    tb_b = s4.shapes.add_textbox(Inches(1.15), Inches(3.90), Inches(11.0), Inches(2.75))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True

    pb1 = tf_b.paragraphs[0]
    pb1.text = "OFFICIAL HANDOUT MANDATE: 'WITH MORE TIME WE'D ADD...'"
    pb1.font.name = "Segoe UI"
    pb1.font.size = Pt(12)
    pb1.font.bold = True
    pb1.font.color.rgb = BLUE_ACCENT

    pb2 = tf_b.add_paragraph()
    pb2.text = "\"With more time, we’d add live Canvas and syllabus LMS integration, automatically syncing weekly exam dates and generating proactive 15-minute micro-sprints before every morning class.\n\nWe don't just help students pass exams; we eliminate academic paralysis and unlock confidence.\""
    pb2.font.name = "Segoe UI"
    pb2.font.size = Pt(13)
    pb2.font.italic = True
    pb2.font.color.rgb = TEXT_WHITE
    pb2.space_before = Pt(6)

    pb3 = tf_b.add_paragraph()
    pb3.text = "WE ARE PRISMATIC — REFRACTING COMPLEX KNOWLEDGE INTO YOUR NATIVE SPECTRUM. VOTE PRISMATIC!"
    pb3.font.name = "Segoe UI"
    pb3.font.size = Pt(12.5)
    pb3.font.bold = True
    pb3.font.color.rgb = CYAN_ACCENT
    pb3.space_before = Pt(8)

    add_slide_footer(s4, 4, 4)

    notes4 = s4.notes_slide.notes_text_frame
    notes4.text = (
        "OFFICIAL TIMED SCRIPT (1:45 - 2:00):\n\n"
        "[1:45 - 2:00] SEGMENT 5: THE CLOSE (15s)\n"
        "\"With more time, we’d add live Canvas and syllabus LMS integration, automatically syncing weekly exam dates and generating proactive 15-minute micro-sprints before every morning class.\n"
        "We are PRISMATIC—refracting complex knowledge into your native spectrum.\n"
        "Vote Prismatic for your room champion! Thank you!\""
    )

    prs.save(str(OUTPUT_PPTX))
    print(f"[+] Successfully generated Executive Slate PPTX: {OUTPUT_PPTX}")

def export_to_pdf_and_images(skip_pdf=False, skip_images=False):
    if skip_pdf and skip_images:
        return

    com_success = False
    try:
        import win32com.client
        print("[*] Attempting export via PowerPoint COM...")
        powerpoint = win32com.client.Dispatch("PowerPoint.Application")
        deck = powerpoint.Presentations.Open(str(OUTPUT_PPTX), WithWindow=False)
        
        try:
            if not skip_pdf:
                deck.SaveAs(str(OUTPUT_PDF), 32)
                print(f"[+] Successfully exported Executive Slate PDF: {OUTPUT_PDF}")

            if not skip_images:
                SLIDES_DIR.mkdir(parents=True, exist_ok=True)
                for i in range(1, deck.Slides.Count + 1):
                    img_path = SLIDES_DIR / f"slide_{i}.png"
                    deck.Slides(i).Export(str(img_path), "PNG", 1920, 1080)
                    print(f"[+] Exported Slide {i} image: {img_path}")
            com_success = True
        finally:
            deck.Close()
            powerpoint.Quit()

    except Exception as e:
        print(f"[!] PowerPoint COM export unavailable or failed ({e}).")

    if not com_success and not skip_pdf:
        print("[*] Attempting PDF conversion via LibreOffice headless (soffice)...")
        try:
            res = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", str(OUTPUT_PPTX), "--outdir", str(BASE_DIR)],
                                 capture_output=True, text=True, check=True)
            print(f"[+] LibreOffice exported PDF successfully.")
        except Exception as err:
            print(f"[!] LibreOffice fallback also unavailable ({err}). PPTX is intact.")

def main():
    parser = argparse.ArgumentParser(description="Prismatic Pitch Deck Builder")
    parser.add_argument("--skip-pdf", action="store_true", help="Skip PDF export")
    parser.add_argument("--skip-images", action="store_true", help="Skip PNG slide image export")
    parser.add_argument("--pptx-only", action="store_true", help="Only build PPTX, skip exports")
    parser.add_argument("--pdf-only", action="store_true", help="Only export PDF from existing PPTX")
    parser.add_argument("--images-only", action="store_true", help="Only export slide images from existing PPTX")
    args = parser.parse_args()

    if args.pdf_only:
        export_to_pdf_and_images(skip_pdf=False, skip_images=True)
        return
    if args.images_only:
        export_to_pdf_and_images(skip_pdf=True, skip_images=False)
        return

    build_prismatic_masterpiece()
    skip_pdf = args.skip_pdf or args.pptx_only
    skip_images = args.skip_images or args.pptx_only
    export_to_pdf_and_images(skip_pdf=skip_pdf, skip_images=skip_images)

if __name__ == "__main__":
    main()
